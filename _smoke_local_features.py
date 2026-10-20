#!/usr/bin/env python3
"""本地验证 PDF 拆分、PDF 转 PPT、PDF 转 Excel 三个新功能。"""
from pathlib import Path
import sys
sys.path.insert(0, str(Path(__file__).parent))

import app

samples = list(Path(__file__).parent.glob('示例*.pdf'))
print('=== 样例 PDF ===')
for s in samples:
    print(f'  {s.name} ({s.stat().st_size / 1024:.1f} KB)')

if not samples:
    print('NO SAMPLES')
    sys.exit(1)

data = samples[0].read_bytes()
doc = app.pymupdf.open(samples[0])
pages = doc.page_count
doc.close()
print(f'\n=== 使用 {samples[0].name} ({pages} 页) ===')

# 1. 拆分 - 按页
print('\n--- 拆分: 按页 ---')
o = app.split_pdf(data, samples[0].name)
print(f'  ok={o.ok}  msg={o.message}')
print(f'  data={len(o.data or b"")} bytes, files={len(o.files or [])}')
if o.files:
    for name, blob in o.files[:3]:
        print(f'    {name} ({len(blob)} bytes)')

# 2. 拆分 - 区间
print('\n--- 拆分: 区间 ---')
o = app.split_pdf(data, samples[0].name, f'1-{pages}')
print(f'  ok={o.ok}  msg={o.message}')

# 3. PPT
print('\n--- PDF -> PPT ---')
o = app.pdf_to_ppt(data, samples[0].name, dpi=72)
print(f'  ok={o.ok}  msg={o.message}')
print(f'  data={len(o.data or b"")} bytes')

# 4. Excel (样例可能没表格)
print('\n--- PDF -> Excel ---')
o = app.pdf_to_excel(data, samples[0].name)
print(f'  ok={o.ok}  msg={o.message}')
if o.data:
    print(f'  data={len(o.data)} bytes')

# 5. 用 reportlab 造一个带表格的 PDF 再测 Excel
print('\n--- PDF -> Excel (人造表格) ---')
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from io import BytesIO
buf = BytesIO()
c = canvas.Canvas(buf, pagesize=A4)
w, h = A4
c.drawString(50, h - 50, 'Test Table')
x, y = 50, h - 80
col_w, row_h = 150, 30
for row in range(3):
    for col in range(3):
        c.rect(x + col * col_w, y - row * row_h, col_w, row_h)
        c.drawString(x + col * col_w + 10, y - row * row_h + 10, f'R{row}C{col}')
c.showPage()
c.save()
table_bytes = buf.getvalue()
o = app.pdf_to_excel(table_bytes, 'test_table.pdf')
print(f'  ok={o.ok}  msg={o.message}')
if o.data:
    print(f'  data={len(o.data)} bytes')
    # 验证 Excel 可读
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(o.data))
    print(f'  sheets={wb.sheetnames}')
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    print(f'  rows={len(rows)}  first_row={rows[0] if rows else None}')

# 6. 边界情况
print('\n--- 边界: 空数据 ---')
o = app.split_pdf(b'', 'empty.pdf')
print(f'  split ok={o.ok}  msg={o.message}')
o = app.pdf_to_ppt(b'', 'empty.pdf')
print(f'  ppt ok={o.ok}  msg={o.message}')
o = app.pdf_to_excel(b'', 'empty.pdf')
print(f'  excel ok={o.ok}  msg={o.message}')

print('\n--- 边界: 非法页码 ---')
o = app.split_pdf(data, samples[0].name, 'abc')
print(f'  ok={o.ok}  msg={o.message}')

print('\n=== ALL DONE ===')

