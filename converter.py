"""PDF -> Word (.docx) conversion core.

Primary engine: pdf2docx (layout-aware conversion).
Fallback engine: PyMuPDF text extraction + python-docx (useful for scanned or
partially broken PDFs, and as a safety net when pdf2docx fails/returns blank).
"""

from __future__ import annotations

import re
import traceback
import warnings
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Iterable, Sequence

# PyMuPDF：优先使用新模块名 pymupdf，兼容旧版别名 fitz。
try:  # pragma: no cover - import shim
    import pymupdf as fitz  # type: ignore
    fitz.TOOLS  # noqa: B018 - touch attribute to ensure the module is real
except (ImportError, AttributeError):  # pragma: no cover
    import fitz  # type: ignore

# 屏蔽 pdf2docx 内部的 fitz 弃用提示（其依赖仍使用旧别名）。
warnings.filterwarnings("ignore", message=r"The `fitz` API is deprecated.*")
warnings.filterwarnings("ignore", category=DeprecationWarning, module=r"pdf2docx.*")

from docx import Document
from docx.shared import Pt

import fonts

ProgressCallback = Callable[[str], None]


@dataclass
class ConversionResult:
    source: Path
    output: Path | None = None
    ok: bool = False
    engine: str = ""
    message: str = ""
    traceback_text: str = ""
    page_count: int = 0
    font_report: fonts.FontReport | None = None


@dataclass
class BatchReport:
    results: list[ConversionResult] = field(default_factory=list)

    @property
    def succeeded(self) -> int:
        return sum(1 for r in self.results if r.ok)

    @property
    def failed(self) -> int:
        return sum(1 for r in self.results if not r.ok)


def _log(callback: ProgressCallback | None, message: str) -> None:
    if callback is not None:
        callback(message)


# --------------------------------------------------------------------------
# Engines
# --------------------------------------------------------------------------
def _convert_with_pdf2docx(
    src: Path,
    dst: Path,
    start_page: int,
    end_page: int,
    callback: ProgressCallback | None,
) -> None:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        from pdf2docx import Converter

    progress = (
        (lambda done, total: _log(callback, f"  正在解析第 {done}/{total} 页"))
        if callback is not None
        else None
    )
    converter = Converter(str(src))
    try:
        # 注意：pdf2docx 的 end 是开区间（内部 slice(start, end)），
        # 因此传 end_page + 1，否则每次都会漏掉最后一页（3 页只转出 2 页）。
        converter.convert(
            str(dst),
            start=start_page,
            end=end_page + 1,
            multi_processing=False,
            progress_callback=progress,
        )
    finally:
        converter.close()


def _convert_with_pymupdf(
    src: Path,
    dst: Path,
    start_page: int,
    end_page: int,
    callback: ProgressCallback | None,
) -> None:
    """Text-only fallback: extract text page by page and write a Word file."""
    document = fitz.open(str(src))
    try:
        last_page = min(end_page, document.page_count - 1)
        docx = Document()
        style = docx.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)

        any_text = False
        for index in range(start_page, last_page + 1):
            page = document.load_page(index)
            text = page.get_text("text").strip()
            _log(callback, f"  正在提取第 {index + 1}/{last_page + 1} 页文字")
            if index > start_page:
                docx.add_page_break()
            if text:
                any_text = True
                for line in text.splitlines():
                    if line.strip():
                        docx.add_paragraph(line)
            else:
                docx.add_paragraph(
                    f"[第 {index + 1} 页：未检测到可提取文字，可能是扫描件，请先做 OCR]"
                )

        if not any_text:
            raise RuntimeError("PDF 中没有可提取的文字（可能是扫描件），建议先做 OCR")
        docx.save(str(dst))
    finally:
        document.close()


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------
def _docx_has_content(path: Path) -> bool:
    """Guard against pdf2docx producing an empty/blank document."""
    try:
        document = Document(str(path))
    except Exception:
        return False
    if any(p.text.strip() for p in document.paragraphs):
        return True
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    return True
    return False


def _docx_page_count(path: Path) -> int | None:
    """统计 docx 中的显式分页符数量 +1，用于确认页数是否被截断。"""
    try:
        with zipfile.ZipFile(path) as archive:
            xml = archive.read("word/document.xml").decode("utf-8", "ignore")
    except (KeyError, zipfile.BadZipFile, OSError):
        return None
    breaks = len(re.findall(r'<w:br[^>]*w:type="page"', xml))
    sections = len(re.findall(r"<w:sectPr", xml))
    return max(breaks + 1, sections)


def _page_range(page_count: int, page_start: int | None, page_end: int | None) -> tuple[int, int]:
    """Convert 1-based user page numbers to a 0-based inclusive range."""
    start = max((page_start or 1) - 1, 0)
    if page_end is None:
        end = page_count - 1
    else:
        end = min(max(page_end - 1, start), page_count - 1)
    return start, end


def unique_output_path(directory: Path, stem: str, suffix: str = ".docx") -> Path:
    """Avoid overwriting: report.docx -> report(1).docx -> report(2).docx ..."""
    candidate = directory / f"{stem}{suffix}"
    counter = 1
    while candidate.exists():
        candidate = directory / f"{stem}({counter}){suffix}"
        counter += 1
    return candidate


def collect_pdfs(paths: Iterable[str | Path], recursive: bool = True) -> list[Path]:
    """Expand files and folders into a de-duplicated, sorted list of PDF paths."""
    found: list[Path] = []
    seen: set[Path] = set()

    for raw in paths:
        path = Path(raw).expanduser()
        if path.is_dir():
            iterator = path.rglob("*") if recursive else path.glob("*")
            candidates = [p for p in iterator if p.is_file() and p.suffix.lower() == ".pdf"]
        elif path.is_file() and path.suffix.lower() == ".pdf":
            candidates = [path]
        else:
            continue
        for candidate in sorted(candidates):
            resolved = candidate.resolve()
            if resolved not in seen:
                seen.add(resolved)
                found.append(candidate)
    return found


# --------------------------------------------------------------------------
# Public API
# --------------------------------------------------------------------------
def convert_pdf(
    src: str | Path,
    out_dir: str | Path | None = None,
    page_start: int | None = None,
    page_end: int | None = None,
    overwrite: bool = False,
    callback: ProgressCallback | None = None,
    *,
    detect_fonts: bool = True,
    auto_download_fonts: bool = False,
    extract_embedded_fonts: bool = True,
) -> ConversionResult:
    """Convert a single PDF to .docx, with automatic fallback.

    Args:
        detect_fonts: 识别 PDF 使用的字体并应用到输出文档的默认样式。
        auto_download_fonts: 本机缺失字体时自动下载开源替代字体（会联网）。
        extract_embedded_fonts: 是否把 PDF 内嵌字体抽到本机缓存目录复用。
            设为 False 可确保不把用户文档的字体子集落盘（「文件不留存」）。
    """
    src_path = Path(src).expanduser().resolve()
    result = ConversionResult(source=src_path)

    if not src_path.is_file():
        result.message = "文件不存在"
        return result

    target_dir = Path(out_dir).expanduser().resolve() if out_dir else src_path.parent
    target_dir.mkdir(parents=True, exist_ok=True)
    dst = target_dir / f"{src_path.stem}.docx"
    if overwrite:
        dst.unlink(missing_ok=True)
    elif dst.exists():
        dst = unique_output_path(target_dir, src_path.stem)
    result.output = dst

    with fitz.open(str(src_path)) as probe:
        page_count = probe.page_count
    start, end = _page_range(page_count, page_start, page_end)
    wanted_pages = end - start + 1
    result.page_count = wanted_pages

    font_report: fonts.FontReport | None = None
    if detect_fonts:
        _log(callback, "识别字体…")
        try:
            records = fonts.scan_pdf_fonts(src_path)
            font_report = fonts.resolve_fonts(
                records, src_path,
                auto_download=auto_download_fonts,
                callback=callback,
                extract_embedded=extract_embedded_fonts,
            )
            result.font_report = font_report
            for record in records:
                resolved = font_report.resolved.get(record.pdf_name)
                if resolved is not None:
                    _log(callback, f"  字体 {record.display} → {resolved.family}（{resolved.source}）")
        except Exception as exc:  # noqa: BLE001 - 字体问题不应中断转换
            _log(callback, f"  字体识别失败（已忽略）：{exc}")

    _log(callback, f"[1/2] pdf2docx 转换：{src_path.name}")
    try:
        _convert_with_pdf2docx(src_path, dst, start, end, callback)
        if not _docx_has_content(dst):
            raise RuntimeError("pdf2docx 输出为空（可能是扫描件或纯图片 PDF）")
        result.ok = True
        result.engine = "pdf2docx"
        result.message = f"完成，共 {wanted_pages} 页"
        _log(callback, f"  ✔ {dst.name}（pdf2docx）")
    except Exception as exc:  # noqa: BLE001 - intentionally fall through
        result.traceback_text += traceback.format_exc()
        _log(callback, f"  pdf2docx 失败：{exc}")
        dst.unlink(missing_ok=True)

        _log(callback, f"[2/2] 回退为文字提取模式：{src_path.name}")
        try:
            _convert_with_pymupdf(src_path, dst, start, end, callback)
            result.ok = True
            result.engine = "pymupdf-text"
            result.message = f"完成（文字模式），共 {wanted_pages} 页"
            _log(callback, f"  ✔ {dst.name}（文字模式）")
        except Exception as fallback_exc:  # noqa: BLE001
            result.traceback_text += traceback.format_exc()
            result.message = f"失败：{fallback_exc}"
            _log(callback, f"  ✘ {src_path.name} 转换失败：{fallback_exc}")
            dst.unlink(missing_ok=True)
            return result

    # 页数校验：确保没有丢页
    try:
        actual = _docx_page_count(dst)
    except Exception:  # noqa: BLE001
        actual = None
    if actual is not None and actual != wanted_pages:
        _log(callback, f"  ⚠ 页数校验：PDF {wanted_pages} 页，Word 检测到 {actual} 页（Word 自动分页可能与 PDF 不同）")

    if font_report is not None and font_report.records:
        try:
            fonts.apply_font_report(dst, font_report)
            result.message += f"；{font_report.summary()}"
            _log(callback, f"  字体：{font_report.summary()}")
        except Exception as exc:  # noqa: BLE001
            _log(callback, f"  应用字体失败（已忽略）：{exc}")
    return result


def convert_many(
    sources: Sequence[str | Path],
    out_dir: str | Path | None = None,
    overwrite: bool = False,
    callback: ProgressCallback | None = None,
) -> BatchReport:
    """Convert several PDFs; per-file output goes into ``out_dir`` when given."""
    report = BatchReport()
    total = len(sources)
    for index, src in enumerate(sources, start=1):
        _log(callback, f"({index}/{total}) {Path(src).name}")
        report.results.append(
            convert_pdf(src, out_dir=out_dir, overwrite=overwrite, callback=callback)
        )
    return report
